"""Deep local document processing for PDFs, Office archives, and export bundles."""

from __future__ import annotations

import base64
import binascii
import difflib
import hashlib
import hmac
import json
import re
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import uuid4
from xml.etree import ElementTree

from .documents import DocumentManager
from .knowledge import KnowledgeEngine
from .models import AuditEvent, utc_now
from .storage import OpenAgentStore


class DocumentProcessing:
    """Converts local files into versioned OpenAgent documents."""

    def __init__(self, store: OpenAgentStore, *, knowledge: KnowledgeEngine | None = None) -> None:
        self.store = store
        self.documents = DocumentManager(store)
        self.knowledge = knowledge or KnowledgeEngine(store)
        self.artifact_dir = self.store.data_dir / "document_processing"
        self.artifact_dir.mkdir(parents=True, exist_ok=True)
        self.initialize()

    def initialize(self) -> None:
        self.store.initialize()

    def import_pdf(
        self,
        *,
        filename: str,
        content: bytes,
        project: str = "default",
        owner: str = "local",
        title: str = "",
    ) -> dict[str, Any]:
        text = extract_pdf_text(content)
        if not text.strip():
            text = f"PDF imported: {Path(filename).name}. No extractable text found with the local fallback extractor."
        doc = self.documents.create_document(
            title=title or Path(filename).stem or "Imported PDF",
            content=text,
            project=project,
            owner=owner,
            doc_format="markdown",
            tags=["pdf", "imported"],
            folder="imports/pdf",
        )
        artifact = self._write_artifact(filename, content, project=project)
        self._register_artifact(doc["doc_id"], artifact, project=project, kind="source-pdf")
        self.store.add_audit_event(
            AuditEvent(
                actor="document-processing",
                action="document.import_pdf",
                resource=f"document:{doc['doc_id']}",
                project=project,
                metadata={"filename": filename, "artifact": str(artifact), "chars": len(text)},
            )
        )
        indexed = self.knowledge.ingest_text(
            title=filename,
            content=text,
            project=project,
            source="document-import",
            source_path=filename,
            metadata={"editor_document_id": doc["doc_id"], "format": "pdf"},
        )
        return {"document": doc, "chars_extracted": len(text), "artifact_path": str(artifact), "knowledge_document": indexed}

    def import_office_or_text(
        self,
        *,
        filename: str,
        content: bytes,
        project: str = "default",
        owner: str = "local",
        title: str = "",
    ) -> dict[str, Any]:
        extension = Path(filename).suffix.lower()
        if extension in {".doc", ".ppt", ".xls"}:
            modern = {".doc": ".docx", ".ppt": ".pptx", ".xls": ".xlsx"}[extension]
            raise ValueError(
                f"Legacy {extension[1:].upper()} files are not supported for text extraction. "
                f"Save the file as {modern} or export it as PDF, then import it again."
            )
        text = extract_office_text(filename, content)
        if not text.strip():
            if extension in {".odt", ".rtf"}:
                raise ValueError(f"Could not extract readable text from this {extension[1:].upper()} file.")
            text = content.decode("utf-8", errors="replace")
        doc = self.documents.create_document(
            title=title or Path(filename).stem or "Imported document",
            content=text,
            project=project,
            owner=owner,
            doc_format="markdown",
            tags=["imported"],
            folder="imports/documents",
        )
        artifact = self._write_artifact(filename, content, project=project)
        self._register_artifact(doc["doc_id"], artifact, project=project, kind="source-document")
        self.store.add_audit_event(
            AuditEvent(
                actor="document-processing",
                action="document.import",
                resource=f"document:{doc['doc_id']}",
                project=project,
                metadata={"filename": filename, "artifact": str(artifact), "chars": len(text)},
            )
        )
        indexed = self.knowledge.ingest_text(
            title=filename,
            content=text,
            project=project,
            source="document-import",
            source_path=filename,
            metadata={"editor_document_id": doc["doc_id"], "format": Path(filename).suffix.lower().lstrip(".") or "text"},
        )
        return {"document": doc, "chars_extracted": len(text), "artifact_path": str(artifact), "knowledge_document": indexed}

    def extract_document_text(self, doc_id: str, *, owner: str = "") -> dict[str, Any]:
        doc = self.documents.get_document(doc_id, owner=owner)
        content = str(doc.get("content", ""))
        return {
            "doc_id": doc_id,
            "title": doc["title"],
            "chars": len(content),
            "text": content,
            "privacy_boundary": "Document text is read from local SQLite editor storage.",
        }

    def render_preview(self, doc_id: str, *, owner: str = "", max_lines: int = 12) -> dict[str, Any]:
        doc = self.documents.get_document(doc_id, owner=owner)
        lines = str(doc.get("content", "")).splitlines()[: max(1, min(max_lines, 100))]
        return {
            "doc_id": doc_id,
            "title": doc["title"],
            "version": doc["version"],
            "preview_lines": lines,
            "line_count": len(str(doc.get("content", "")).splitlines()),
            "render_mode": "local-text-preview",
        }

    def render_pdf_page(self, doc_id: str, *, page: int = 0, scale: float = 1.5, owner: str = "") -> dict[str, Any]:
        self.documents.get_document(doc_id, owner=owner)
        source = self._latest_artifact(doc_id, kind="source-pdf")
        if source is None:
            raise ValueError("Document has no imported PDF artifact.")
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(source["local_path"])
        page_object = None
        bitmap = None
        try:
            page_count = len(pdf)
            if page < 0 or page >= page_count:
                raise ValueError("PDF page is out of range.")
            page_object = pdf[page]
            bitmap = page_object.render(scale=max(0.5, min(float(scale), 4.0)))
            output = self.artifact_dir / f"{doc_id}-page-{page + 1}-{uuid4().hex[:8]}.png"
            bitmap.to_pil().save(output, format="PNG", optimize=True)
        finally:
            if bitmap is not None:
                bitmap.close()
            if page_object is not None:
                page_object.close()
            pdf.close()
        artifact = self._register_artifact(doc_id, output, project=source["project"], kind="page-render", metadata={"page": page})
        return {"doc_id": doc_id, "page": page, "page_count": page_count, "artifact": artifact}

    def add_annotation(
        self,
        doc_id: str,
        *,
        page: int,
        kind: str = "note",
        geometry: dict[str, Any] | None = None,
        text: str = "",
        metadata: dict[str, Any] | None = None,
        owner: str = "",
    ) -> dict[str, Any]:
        self.documents.get_document(doc_id, owner=owner)
        annotation_id = uuid4().hex
        with self.store.session() as connection:
            connection.execute(
                """
                INSERT INTO document_annotations (id, doc_id, page, kind, geometry_json, text, metadata_json, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (annotation_id, doc_id, max(0, int(page)), kind, json.dumps(geometry or {}, sort_keys=True), text, json.dumps(metadata or {}, sort_keys=True), utc_now()),
            )
        return {"id": annotation_id, "doc_id": doc_id, "page": max(0, int(page)), "kind": kind, "geometry": geometry or {}, "text": text, "metadata": metadata or {}}

    def annotations(self, doc_id: str, *, owner: str = "") -> list[dict[str, Any]]:
        self.documents.get_document(doc_id, owner=owner)
        with self.store.session() as connection:
            rows = connection.execute("SELECT * FROM document_annotations WHERE doc_id = ? ORDER BY page, created_at", (doc_id,)).fetchall()
        return [
            {"id": row["id"], "doc_id": row["doc_id"], "page": row["page"], "kind": row["kind"], "geometry": json.loads(row["geometry_json"]), "text": row["text"], "metadata": json.loads(row["metadata_json"]), "created_at": row["created_at"]}
            for row in rows
        ]

    def add_ocr_overlay(self, doc_id: str, *, page: int, text: str, blocks: list[dict[str, Any]] | None = None, owner: str = "") -> dict[str, Any]:
        return self.add_annotation(doc_id, page=page, kind="ocr-overlay", text=text, metadata={"blocks": blocks or [], "engine": "local-submitted"}, owner=owner)

    def redact_text(self, doc_id: str, patterns: list[str], *, replacement: str = "[REDACTED]", owner: str = "") -> dict[str, Any]:
        document = self.documents.get_document(doc_id, owner=owner)
        content = str(document.get("content", ""))
        replacements = 0
        for pattern in patterns:
            if not pattern:
                continue
            content, count = re.subn(re.escape(pattern), replacement, content, flags=re.IGNORECASE)
            replacements += count
        updated = self.documents.update_document(doc_id, content=content, if_version=document["version"], owner=owner)
        self.store.add_audit_event(
            AuditEvent(actor="document-processing", action="document.redact", resource=f"document:{doc_id}", project=document["project"], metadata={"patterns": len(patterns), "replacements": replacements})
        )
        return {"document": updated, "replacements": replacements, "redaction_mode": "versioned-text-layer"}

    def export_pdf(self, doc_id: str, *, owner: str = "", signing_key: str = "") -> dict[str, Any]:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen.canvas import Canvas

        document = self.documents.get_document(doc_id, owner=owner)
        path = self.artifact_dir / f"{doc_id}-v{document['version']}.pdf"
        canvas = Canvas(str(path), pagesize=A4, invariant=1, pageCompression=1)
        width, height = A4
        text = canvas.beginText(42, height - 54)
        text.setFont("Helvetica", 10)
        for line in str(document.get("content", "")).splitlines() or [""]:
            for chunk in [line[index:index + 95] for index in range(0, max(1, len(line)), 95)]:
                if text.getY() < 54:
                    canvas.drawText(text)
                    canvas.showPage()
                    text = canvas.beginText(42, height - 54)
                    text.setFont("Helvetica", 10)
                text.textLine(chunk)
        canvas.drawText(text)
        canvas.save()
        artifact = self._register_artifact(doc_id, path, project=document["project"], kind="pdf-export", metadata={"version": document["version"], "deterministic": True})
        signature = hmac.new(signing_key.encode("utf-8"), path.read_bytes(), hashlib.sha256).hexdigest() if signing_key else ""
        return {"document": document, "artifact": artifact, "signature": signature, "signature_algorithm": "HMAC-SHA256" if signature else "", "deterministic": True}

    def export_docx(self, doc_id: str, *, owner: str = "") -> dict[str, Any]:
        try:
            from docx import Document
        except ImportError as exc:
            raise ValueError("DOCX export requires `pip install openagent[office]`.") from exc
        document = self.documents.get_document(doc_id, owner=owner)
        path = self.artifact_dir / f"{doc_id}-v{document['version']}.docx"
        output = Document()
        output.add_heading(str(document["title"]), level=1)
        for paragraph in str(document.get("content", "")).split("\n\n"):
            output.add_paragraph(paragraph)
        output.save(path)
        artifact = self._register_artifact(doc_id, path, project=document["project"], kind="docx-export", metadata={"version": document["version"]})
        self._audit_export(document, "docx", path)
        return {"document": document, "artifact": artifact, "format": "docx"}

    def export_pptx(self, doc_id: str, *, owner: str = "") -> dict[str, Any]:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise ValueError("PPTX export requires `pip install openagent[office]`.") from exc
        document = self.documents.get_document(doc_id, owner=owner)
        path = self.artifact_dir / f"{doc_id}-v{document['version']}.pptx"
        presentation = Presentation()
        for index, paragraph in enumerate(str(document.get("content", "")).split("\n\n")):
            slide = presentation.slides.add_slide(presentation.slide_layouts[0 if index == 0 else 1])
            slide.shapes.title.text = str(document["title"]) if index == 0 else f"{document['title']} ({index + 1})"
            if len(slide.placeholders) > 1:
                slide.placeholders[1].text = paragraph
        presentation.save(path)
        artifact = self._register_artifact(doc_id, path, project=document["project"], kind="pptx-export", metadata={"version": document["version"]})
        self._audit_export(document, "pptx", path)
        return {"document": document, "artifact": artifact, "format": "pptx"}

    def export_xlsx(self, doc_id: str, *, owner: str = "") -> dict[str, Any]:
        try:
            from openpyxl import Workbook
        except ImportError as exc:
            raise ValueError("XLSX export requires `pip install openagent[office]`.") from exc
        document = self.documents.get_document(doc_id, owner=owner)
        path = self.artifact_dir / f"{doc_id}-v{document['version']}.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "OpenAgent Export"
        sheet.append(["Document", str(document["title"])])
        sheet.append(["Version", document["version"]])
        sheet.append([])
        sheet.append(["Line", "Content"])
        for index, line in enumerate(str(document.get("content", "")).splitlines(), start=1):
            sheet.append([index, line])
        workbook.save(path)
        artifact = self._register_artifact(doc_id, path, project=document["project"], kind="xlsx-export", metadata={"version": document["version"]})
        self._audit_export(document, "xlsx", path)
        return {"document": document, "artifact": artifact, "format": "xlsx"}

    def compare(self, left_doc_id: str, right_doc_id: str, *, owner: str = "") -> dict[str, Any]:
        left = self.documents.get_document(left_doc_id, owner=owner)
        right = self.documents.get_document(right_doc_id, owner=owner)
        diff = list(difflib.unified_diff(str(left.get("content", "")).splitlines(), str(right.get("content", "")).splitlines(), fromfile=left["title"], tofile=right["title"], lineterm=""))
        ratio = difflib.SequenceMatcher(None, str(left.get("content", "")), str(right.get("content", ""))).ratio()
        return {"left": {"doc_id": left_doc_id, "version": left["version"]}, "right": {"doc_id": right_doc_id, "version": right["version"]}, "similarity": round(ratio, 6), "diff": diff}

    def export_zip(self, *, doc_ids: list[str], project: str = "default", owner: str = "") -> dict[str, Any]:
        export_id = uuid4().hex
        path = self.artifact_dir / f"document-export-{export_id}.zip"
        exported = []
        with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            manifest = {"project": project, "created_at": utc_now(), "documents": []}
            for doc_id in doc_ids:
                doc = self.documents.get_document(doc_id, owner=owner)
                safe_title = re.sub(r"[^A-Za-z0-9_.-]+", "_", doc["title"]).strip("_") or doc_id
                name = f"{safe_title}-v{doc['version']}.md"
                archive.writestr(name, str(doc.get("content", "")))
                manifest["documents"].append({"doc_id": doc_id, "title": doc["title"], "file": name, "version": doc["version"]})
                exported.append(doc_id)
            archive.writestr("manifest.json", json.dumps(manifest, indent=2, sort_keys=True))
        self.store.add_audit_event(
            AuditEvent(
                actor="document-processing",
                action="document.export_zip",
                resource=f"document_export:{export_id}",
                project=project,
                metadata={"count": len(exported), "path": str(path)},
            )
        )
        return {"path": str(path), "count": len(exported), "documents": exported}

    def decode_base64(self, value: str) -> bytes:
        try:
            return base64.b64decode(value.encode("ascii"), validate=True)
        except (ValueError, binascii.Error) as exc:
            raise ValueError("content_base64 must be valid base64.") from exc

    def _write_artifact(self, filename: str, content: bytes, *, project: str) -> Path:
        folder = self.artifact_dir / project
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / f"{uuid4().hex}-{Path(filename or 'document.bin').name}"
        path.write_bytes(content)
        return path

    def _register_artifact(self, doc_id: str, path: Path, *, project: str, kind: str, metadata: dict[str, Any] | None = None) -> dict[str, Any]:
        artifact_id = uuid4().hex
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        created = utc_now()
        with self.store.session() as connection:
            connection.execute(
                "INSERT INTO document_artifacts (id, doc_id, project, kind, local_path, sha256, metadata_json, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                (artifact_id, doc_id, project, kind, str(path), digest, json.dumps(metadata or {}, sort_keys=True), created),
            )
        return {"id": artifact_id, "doc_id": doc_id, "project": project, "kind": kind, "local_path": str(path), "sha256": digest, "metadata": metadata or {}, "created_at": created}

    def _latest_artifact(self, doc_id: str, *, kind: str) -> dict[str, Any] | None:
        with self.store.session() as connection:
            row = connection.execute("SELECT * FROM document_artifacts WHERE doc_id = ? AND kind = ? ORDER BY rowid DESC LIMIT 1", (doc_id, kind)).fetchone()
        return dict(row) if row else None

    def _audit_export(self, document: dict[str, Any], export_format: str, path: Path) -> None:
        self.store.add_audit_event(
            AuditEvent(
                actor="document-processing",
                action=f"document.export_{export_format}",
                resource=f"document:{document['doc_id']}",
                project=document["project"],
                metadata={"path": str(path), "version": document["version"]},
            )
        )


def extract_pdf_text(content: bytes) -> str:
    try:
        import pypdfium2 as pdfium

        doc = pdfium.PdfDocument(content)
        pages = []
        for i in range(len(doc)):
            textpage = doc[i].get_textpage()
            pages.append(textpage.get_text_range().strip())
        extracted = "\n".join(p for p in pages if p).strip()
        if extracted:
            return extracted
    except Exception:
        pass
    try:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(content))
        extracted = "\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()
        if extracted:
            return extracted
    except Exception:
        pass
    pieces = []
    for block in re.findall(rb"BT(.*?)ET", content, re.DOTALL):
        for text in re.findall(rb"\((.*?)\)", block):
            decoded = text.decode("latin-1", errors="replace").strip()
            if decoded:
                pieces.append(decoded)
    if not pieces:
        for text in re.findall(rb"\((.*?)\)", content):
            decoded = text.decode("latin-1", errors="replace").strip()
            if decoded:
                pieces.append(decoded)
    return " ".join(pieces)


def extract_office_text(filename: str, content: bytes) -> str:
    lowered = filename.lower()
    if lowered.endswith(".odt"):
        try:
            with zipfile.ZipFile(BytesIO(content)) as archive:
                root = ElementTree.fromstring(archive.read("content.xml"))
            paragraphs = []
            for element in root.iter():
                if element.tag.rsplit("}", 1)[-1] in {"p", "h"}:
                    text = "".join(element.itertext()).strip()
                    if text:
                        paragraphs.append(text)
            return "\n\n".join(paragraphs)
        except (KeyError, ElementTree.ParseError, zipfile.BadZipFile):
            return ""
    if lowered.endswith(".rtf"):
        return _extract_rtf_text(content)
    if lowered.endswith(".docx"):
        try:
            from docx import Document as WordDocument

            document = WordDocument(BytesIO(content))
            paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            if paragraphs:
                return "\n\n".join(paragraphs)
        except (ImportError, ValueError, zipfile.BadZipFile):
            pass
    if lowered.endswith((".docx", ".pptx", ".xlsx")):
        try:
            with zipfile.ZipFile(BytesIO(content)) as archive:
                xml_text = []
                for name in archive.namelist():
                    if name.endswith(".xml") and any(part in name for part in ("word/", "ppt/", "xl/")):
                        raw = archive.read(name).decode("utf-8", errors="replace")
                        cleaned = re.sub(r"<[^>]+>", " ", raw)
                        xml_text.append(re.sub(r"\s+", " ", cleaned).strip())
                return "\n".join(part for part in xml_text if part)
        except zipfile.BadZipFile:
            return ""
    return content.decode("utf-8", errors="replace")


def _extract_rtf_text(content: bytes) -> str:
    raw = content.decode("cp1252", errors="replace")
    if not re.match(r"^\s*\{\\rtf", raw):
        return ""
    ignored_groups = {
        "colortbl", "datastore", "filetbl", "fonttbl", "generator", "info",
        "listoverridetable", "listtable", "object", "pict", "stylesheet",
        "xmlopen", "xmlattrname", "xmlattrvalue", "fldinst",
    }
    ignored = [False]
    output: list[str] = []
    index = 0
    while index < len(raw):
        char = raw[index]
        if char == "{":
            probe = index + 1
            while probe < len(raw) and raw[probe].isspace():
                probe += 1
            if raw.startswith(r"\*", probe):
                probe += 2
                while probe < len(raw) and raw[probe].isspace():
                    probe += 1
            destination = re.match(r"\\([A-Za-z]+)", raw[probe:])
            ignored.append(ignored[-1] or bool(destination and destination.group(1).lower() in ignored_groups))
            index += 1
            continue
        if char == "}":
            if len(ignored) > 1:
                ignored.pop()
            index += 1
            continue
        if char != "\\":
            if not ignored[-1] and char not in "\r\n":
                output.append(char)
            index += 1
            continue
        if index + 1 >= len(raw):
            break
        symbol = raw[index + 1]
        if symbol in "\\{}":
            if not ignored[-1]:
                output.append(symbol)
            index += 2
            continue
        if symbol == "'" and index + 3 < len(raw):
            try:
                decoded = bytes([int(raw[index + 2:index + 4], 16)]).decode("cp1252")
            except (ValueError, UnicodeDecodeError):
                decoded = ""
            if not ignored[-1]:
                output.append(decoded)
            index += 4
            continue
        control = re.match(r"\\([A-Za-z]+)(-?\d+)? ?", raw[index:])
        if control:
            word = control.group(1).lower()
            end = index + control.end()
            if not ignored[-1]:
                if word in {"par", "line"}:
                    output.append("\n")
                elif word == "tab":
                    output.append("\t")
                elif word == "u" and control.group(2):
                    codepoint = int(control.group(2))
                    output.append(chr(codepoint if codepoint >= 0 else codepoint + 65536))
                    if end < len(raw) and raw[end] not in "\\{}":
                        end += 1  # Skip the single fallback character used by \uN.
            if word == "bin" and control.group(2):
                end += max(0, int(control.group(2)))
            index = min(end, len(raw))
            continue
        if not ignored[-1] and symbol == "~":
            output.append(" ")
        elif not ignored[-1] and symbol == "_":
            output.append("-")
        index += 2
    text = "".join(output)
    text = re.sub(r"[ \t]{2,}", " ", text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()
